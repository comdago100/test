import os
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
from dataclasses import replace
from datetime import date
from decimal import Decimal as D
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from inventory_ledger.engine import ORDER, calculate, load_and_calculate, material_entries, product_entries, verify
from inventory_ledger.exporter import HEADERS, export_workbook
from inventory_ledger.loader import ExcelLoadError, parse_day, read_records
from inventory_ledger.models import CATEGORIES, ROLES, END, START, Entry, Record


def record(qty, amount, day=START, name="철판", spec="A", unit="", factory=""):
    q, a = D(str(qty)), D(str(amount))
    return Record(name, spec, unit, day, q, a / q if q else D(0), a, factory)


def input_files(directory):
    paths = {}
    for category in CATEGORIES:
        for role, qty, amount in [("이월", 10, 100), ("거래", 20, 200), ("재고", 5, 50)]:
            path = Path(directory) / f"{category}_{role}.xlsx"
            data = {"재질": [" SUS304 "]} if category == "원재료" else {"품명": [" 철판 "], "규격": [" A "]}
            if category == "제품":
                data = {"공장": [" 1공장 "], **data}
            if role == "거래":
                data = {"출고일" if category == "제품" else "입고일": ["2022-06-01"], **data}
            data["수량"] = [qty]
            if category != "제품":
                data["단가"] = [10]
            if category != "제품" or role != "거래":
                data["금액"] = [amount]
            pd.DataFrame(data).to_excel(path, index=False)
            paths[category, role] = str(path)
    return paths


class LedgerTests(unittest.TestCase):
    def test_quantity_and_amount_integer_rounding(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "부재료_입고.xlsx"
            pd.DataFrame({"품명": ["볼트"] * 4, "규격": ["M8"] * 4,
                          "입고일": ["2022-06-01"] * 4,
                          "수량": [" 1,000.49 ", 3.6, "2.5", "2.6"], "단가": [12.345, 10, 10, 10.5],
                          "금액": [" 129,096.0001 ", 129096.9, "12.5", ""]}).to_excel(path, index=False)
            rows = read_records(path, "거래", "부재료")
            self.assertEqual([r.qty for r in rows], [D(1000), D(4), D(2), D(3)])
            self.assertEqual([r.amount for r in rows], [D(129096), D(129097), D(12), D(32)])
            self.assertEqual(rows[0].price, D("12.345"))
            self.assertIn("129,096.0001", rows[0].source.describe())

    def test_rounding_all_nine_files_still_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = input_files(directory)
            for path in paths.values():
                frame = pd.read_excel(path)
                frame["수량"] = frame["수량"].map(lambda v: f"{v + 0.0001:,.4f}")
                if "금액" in frame:
                    frame["금액"] = frame["금액"].map(lambda v: f"{v + 0.0001:,.4f}")
                frame.to_excel(path, index=False)
            result = load_and_calculate(paths)
            self.assertTrue(result.valid, result.issues)
            export_workbook(result, Path(directory) / "output.xlsx", "제조회사")
            self.assertTrue(all(e.record.qty == int(e.record.qty) and e.record.amount == int(e.record.amount)
                                for groups in result.ledgers.values() for entries in groups for e in entries))

    def test_actual_excel_row_137_and_detailed_multiple_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = input_files(directory)
            path = Path(directory) / "부재료_입고.xlsx"
            blank = {"품명": "", "규격": "", "입고일": "", "수량": "", "단가": "", "금액": ""}
            rows = [dict(blank) for _ in range(135)]
            rows += [{"품명": " 볼트 ", "규격": " M8 ", "입고일": "2022-06-01", "수량": "1,000oops", "단가": 10, "금액": 100},
                     dict(blank),
                     {"품명": " 너트 ", "규격": " M10 ", "입고일": "2022-06-02", "수량": -1, "단가": 10, "금액": 100}]
            pd.DataFrame(rows).to_excel(path, index=False)
            paths["부재료", "거래"] = str(path)
            result = load_and_calculate(paths)
            self.assertFalse(result.valid)
            self.assertEqual(len(result.issues), 2)
            first, second = result.issues
            self.assertEqual(first.item, "[입고] 볼트 / M8")
            self.assertIn("엑셀 137행", first.message)
            self.assertIn("부재료_입고.xlsx", first.message)
            self.assertIn("수량='1,000oops'", first.message)
            self.assertIn("품명=' 볼트 '", first.message)
            self.assertIn("규격=' M8 '", first.message)
            self.assertIn("엑셀 139행", second.message)
            self.assertIn("수량=-1", second.message)

    def test_formula_cached_values_are_rounded_and_missing_values_explained(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "원재료_입고.xlsx"
            pd.DataFrame({"재질": ["SUS304"], "입고일": ["2022-06-01"], "수량": ["=2+1.6"],
                          "단가": [10], "금액": ["=129096+0.0001"]}).to_excel(path, index=False)
            # Excel이 저장하는 수식 캐시를 재현한다. 수식과 실제 읽은 숫자를 모두 보존한다.
            with ZipFile(path) as archive:
                contents = {info.filename: archive.read(info.filename) for info in archive.infolist()}
            namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            sheet = ET.fromstring(contents["xl/worksheets/sheet1.xml"])
            for cell in sheet.findall(".//s:c", namespace):
                if cell.attrib["r"] in ("C2", "E2"):
                    cell.find("s:v", namespace).text = "3.6" if cell.attrib["r"] == "C2" else "129096.0001"
            contents["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
            with ZipFile(path, "w", ZIP_DEFLATED) as archive:
                for filename, content in contents.items():
                    archive.writestr(filename, content)
            records = read_records(path, "거래", "원재료")
            self.assertEqual(records[0].qty, D(4))
            self.assertEqual(records[0].amount, D(129096))
            self.assertIn("=129096+0.0001", records[0].source.describe())
            self.assertIn("129096.0001", records[0].source.describe())
            # 캐시 없는 수식을 원가 0원이나 빈 행으로 처리하지 않는다.
            pd.DataFrame({"재질": ["SUS304"], "입고일": ["2022-06-01"], "수량": [4],
                          "단가": [10], "금액": ["=4*10"]}).to_excel(path, index=False)
            with self.assertRaisesRegex(ExcelLoadError, "저장된 계산값") as context:
                read_records(path, "거래", "원재료")
            self.assertIn("=4*10", str(context.exception))
            self.assertIn("엑셀 2행", str(context.exception))

    def test_calculation_errors_include_source_values_and_file_kind(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = input_files(directory)
            path = paths["부재료", "재고"]
            frame = pd.read_excel(path)
            frame["수량"] = "1,000.01"
            frame.to_excel(path, index=False)
            result = load_and_calculate(paths)
            error = next(issue for issue in result.issues if issue.level == "오류")
            self.assertIn("철판 / A", error.item)
            self.assertIn("[이월/입고/재고]", error.item)
            self.assertIn("수량='1,000.01'", error.message)
            self.assertIn("정수화 합계: 수량=1000", error.message)
            self.assertIn("엑셀 2행", error.message)

    def test_exact_nine_schemas_without_unit_or_generic_date(self):
        with tempfile.TemporaryDirectory() as directory:
            paths = input_files(directory)
            for (category, role), path in paths.items():
                rows = read_records(path, role, category)
                self.assertEqual(rows[0].unit, "")
                self.assertEqual(rows[0].name, "SUS304" if category == "원재료" else "철판")
                self.assertEqual(rows[0].day, START if role == "이월" else END if role == "재고" else date(2022, 6, 1))
                if category == "제품":
                    self.assertEqual(rows[0].factory, "1공장")
                    self.assertEqual(rows[0].price, D(0) if role == "거래" else D(10))
            self.assertTrue(load_and_calculate(paths).valid)

    def test_material_moving_average_rounding_business_days(self):
        entries = material_entries(record(3, 10), [record(3, 10, date(2022, 6, 4))], record(2, 7, END))
        verify(entries)
        self.assertEqual(sum(e.record.qty for e in entries if e.kind == "출고"), D(4))
        self.assertEqual(sum(e.record.amount for e in entries if e.kind == "출고"), D(13))
        self.assertTrue(all(e.record.day.weekday() < 5 for e in entries if e.kind == "출고"))
        self.assertEqual([(e.record.day, ORDER[e.kind]) for e in entries], sorted((e.record.day, ORDER[e.kind]) for e in entries))

    def test_year_end_business_day_constraint(self):
        with self.assertRaisesRegex(ValueError, "영업일"):
            material_entries(record(0, 0), [record(10, 100, END)], record(0, 0, END))
        entries = material_entries(record(0, 0), [record(10, 100, END)], record(10, 100, END))
        verify(entries)

    def test_product_quantity_only_price_and_same_day_order(self):
        day = date(2022, 2, 1)
        entries = product_entries(record(2, 20), [record(4, 0, day), record(6, 0, day)], record(3, 30, END))
        verify(entries)
        self.assertEqual(sum(e.record.qty for e in entries if e.kind == "입고"), D(11))
        daily = [e for e in entries if e.record.day == day]
        self.assertEqual([e.kind for e in daily], ["입고", "출고", "출고"])
        self.assertTrue(all(e.record.price == D(10) for e in entries if e.kind in ("입고", "출고")))
        self.assertEqual(sum(e.record.amount for e in entries if e.kind == "출고"), D(100))

    def test_product_opening_price_when_closing_zero(self):
        entries = product_entries(record(10, 100), [record(20, 0, date(2022, 7, 1))], record(0, 0, END))
        verify(entries)
        self.assertTrue(all(e.record.price == D(10) for e in entries if e.kind in ("입고", "출고")))
        self.assertEqual(entries[-1].balance_amount, D(0))

    def test_product_closing_price_precedence_and_amount_identity(self):
        entries = product_entries(record(10, 100), [record(5, 0, date(2022, 6, 1))], record(10, 200, END))
        verify(entries)
        self.assertTrue(all(e.record.price == D(20) for e in entries if e.kind in ("입고", "출고")))
        self.assertEqual(entries[-1].balance_amount, D(200))
        self.assertEqual([e.record.amount for e in entries if e.kind == "입고"], [D(200)])

    def test_same_product_in_different_factories_stays_separate(self):
        data = {(c, r): [] for c in CATEGORIES for r in ROLES}
        for factory, price in [("1공장", 10), ("2공장", 20)]:
            data["제품", "이월"].append(record(10, 10 * price, factory=factory))
            data["제품", "거래"].append(record(20, 0, date(2022, 6, 1), factory=factory))
            data["제품", "재고"].append(record(5, 5 * price, END, factory=factory))
        result = calculate(data)
        self.assertTrue(result.valid, result.issues)
        self.assertEqual(len(result.ledgers["제품"]), 2)
        self.assertEqual([es[-1].balance_amount for es in result.ledgers["제품"]], [D(50), D(100)])
        self.assertEqual(record(1, 1, spec="ignored").item_key("원재료"), ("철판",))

    def test_quantity_amount_inversion_and_incompatible_valuation(self):
        with self.assertRaisesRegex(ValueError, "수량 역전"):
            material_entries(record(2, 20), [], record(3, 30, END))
        with self.assertRaisesRegex(ValueError, "금액 역전"):
            material_entries(record(2, 20), [], record(1, 30, END))
        with self.assertRaisesRegex(ValueError, "이동평균"):
            material_entries(record(10, 100), [], record(5, 20, END))
        with self.assertRaisesRegex(ValueError, "수량 역전"):
            product_entries(record(10, 100), [], record(5, 50, END))
        with self.assertRaises(ValueError):
            material_entries(record(10, 100), [], record(10, 90, END))

    def test_dates_header_text_strip_and_validation(self):
        self.assertEqual(parse_day(20220601), date(2022, 6, 1))
        self.assertEqual(parse_day(44713), date(2022, 6, 1))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            frame = pd.DataFrame({" 재질 ": [" NA "], " 입고일 ": [" 2022-06-01 "],
                                  "수량": [" 1,000 "], "단가": [10], "금액": [10000]})
            frame.to_excel(path, index=False)
            rows = read_records(path, "거래", "원재료")
            self.assertEqual(rows[0].item_key("원재료"), ("NA",))
            self.assertEqual(rows[0].amount, D(10000))
            frame[" 입고일 "] = "2023-01-01"
            frame.to_excel(path, index=False)
            with self.assertRaisesRegex(ValueError, "2022"):
                read_records(path, "거래", "원재료")
            frame.drop(columns="금액").to_excel(path, index=False)
            with self.assertRaisesRegex(ValueError, "필수 열 누락"):
                read_records(path, "거래", "원재료")

    def test_final_workbook_structure_content_and_print_settings(self):
        with tempfile.TemporaryDirectory() as directory:
            result = load_and_calculate(input_files(directory))
            self.assertTrue(result.valid, result.issues)
            output = Path(directory) / "ledger.xlsx"
            export_workbook(result, output, " 제조회사 ")
            wb = load_workbook(output)
            self.assertEqual(wb.sheetnames, ["수불집계표", "원재료 수불부", "부재료 수불부", "제품 수불부"])
            for ws in wb:
                self.assertEqual(ws["A1"].value, "2022년도 재고수불부")
                self.assertEqual(ws["A1"].font.size, 18)
                self.assertEqual(ws["A1"].font.underline, "single")
                self.assertEqual(ws["A2"].value, "회사명 : 제조회사")
                self.assertEqual(ws.page_setup.orientation, "landscape")
                self.assertEqual(str(ws.page_setup.paperSize), str(ws.PAPERSIZE_A4))
                self.assertEqual(ws.page_setup.fitToWidth, 1)
                self.assertEqual(ws.page_setup.fitToHeight, 0)
                self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
                self.assertEqual(ws.print_title_rows, "$1:$6")
                self.assertEqual(ws.page_margins.left, 0.3)
                self.assertEqual(ws.oddFooter.center.text, "- &P / &N -")
                for row in ws:
                    for cell in row:
                        if isinstance(cell.value, str):
                            self.assertFalse(any(w in cell.value for w in ("역산", "추정", "자동생성", "보정")))
                if ws.title == "수불집계표":
                    continue
                category = ws.title.split()[0]
                self.assertEqual([cell.value for cell in ws[6]], HEADERS[category])
                self.assertEqual(ws.max_column, {"원재료": 12, "부재료": 13, "제품": 14}[category])
                self.assertEqual(ws["A7"].number_format, "yyyy-mm-dd")
                self.assertEqual(ws.cell(7, ws.max_column).number_format, "#,##0")
                self.assertTrue(set(ws.cell(i, 2).value for i in range(7, ws.max_row)).issubset({"기초이월", "입고", "출고", "기말재고"}))
                self.assertEqual(ws.cell(ws.max_row, 3).value, "[ 소계 ]")
                self.assertEqual(ws.cell(ws.max_row, 3).fill.fgColor.rgb, "00EEEEEE")
            wb.close()

    def test_approval_optional_text_formula_safety_and_atomic_guard(self):
        with tempfile.TemporaryDirectory() as directory:
            result = load_and_calculate(input_files(directory))
            output = Path(directory) / "ledger.xlsx"
            export_workbook(result, output, "=1+1", approval=False)
            wb = load_workbook(output)
            for ws in wb:
                self.assertFalse(any(cell.value == "담당" for row in ws for cell in row))
                self.assertFalse(any(cell.data_type == "f" for row in ws for cell in row))
            wb.close()
            original = output.read_bytes()
            with self.assertRaises(ValueError):
                export_workbook(result, output, "추정회사")
            self.assertEqual(output.read_bytes(), original)

    def test_mismatches_missing_files_and_negative_inventory_simulation(self):
        data = {(c, r): [] for c in CATEGORIES for r in ROLES}
        data["원재료", "거래"] = [record(10, 100)]
        data["원재료", "재고"] = [record(5, 50, END)]
        result = calculate(data)
        self.assertTrue(result.valid)
        self.assertTrue(any("품목 미일치" in i.message for i in result.issues))
        self.assertFalse(load_and_calculate({}).valid)
        with self.assertRaisesRegex(ValueError, "음수"):
            verify([Entry(record(1, 10), "기초이월"), Entry(record(2, 20, date(2022, 6, 1)), "출고")])

    def test_gui_background_validation_progress_and_invalidation(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtCore import QEventLoop, QTimer
        from PySide6.QtWidgets import QApplication
        from inventory_ledger.app import MainWindow
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as directory:
            paths = input_files(directory)
            window = MainWindow()
            self.assertEqual(len(window.slots), 9)
            self.assertFalse(window.export_button.isEnabled())
            for key, path in paths.items():
                window.slots[key].setText(path)
            loop = QEventLoop()
            window.start()
            window.thread.finished.connect(loop.quit)
            timer = QTimer()
            timer.setSingleShot(True)
            timer.timeout.connect(loop.quit)
            timer.start(10000)
            loop.exec()
            timer.stop()
            self.assertIsNone(window.thread)
            self.assertEqual(window.table.rowCount(), 1)
            self.assertEqual(window.table.item(0, 0).text(), "정상")
            self.assertTrue(window.export_button.isEnabled())
            self.assertEqual(window.progress.value(), 100)
            window.slots["원재료", "이월"].setText("changed.xlsx")
            self.assertFalse(window.export_button.isEnabled())
            window.close()
            app.processEvents()

    def test_folder_matching_with_real_filename_roles(self):
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication
        from inventory_ledger.app import MainWindow
        app = QApplication.instance() or QApplication([])
        with tempfile.TemporaryDirectory() as directory:
            for c in CATEGORIES:
                for r in ROLES:
                    marker = ("출고" if c == "제품" else "입고") if r == "거래" else r
                    (Path(directory) / f"{c}_{marker}.xlsx").touch()
            window = MainWindow()
            with patch("inventory_ledger.app.QFileDialog.getExistingDirectory", return_value=directory):
                window.load_folder()
            self.assertTrue(all(slot.text() for slot in window.slots.values()))
            (Path(directory) / "제품_출고_중복.xlsx").touch()
            with patch("inventory_ledger.app.QFileDialog.getExistingDirectory", return_value=directory):
                window.load_folder()
            self.assertEqual(window.slots["제품", "거래"].text(), "")
            window.close()
            app.processEvents()


if __name__ == "__main__":
    unittest.main()
