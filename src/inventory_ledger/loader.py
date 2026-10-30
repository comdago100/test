"""입력 파일의 첫 번째 시트, 첫 행을 열 제목으로 읽는다."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl import load_workbook

from .models import END, START, Issue, Record, Source, file_kind

ALIASES = {
    "품명": ("원재료품명", "부재료품명", "제품품명", "품목명"),
    "규격": ("원재료규격", "부재료규격", "제품규격"),
    "단위": ("원재료단위", "부재료단위", "제품단위", "기준단위"),
    "일자": ("입고일", "출고일", "입고일자", "출고일자", "거래일자", "날짜"),
    "수량": ("이월수량", "입고수량", "출고수량", "재고수량"),
    "단가": ("이월단가", "입고단가", "출고단가", "재고단가"),
    "금액": ("이월금액", "입고금액", "출고금액", "재고금액"),
}
NORMAL = {alias: key for key, aliases in ALIASES.items() for alias in (key, *aliases)}


class ExcelLoadError(ValueError):
    def __init__(self, issues: list[Issue]):
        self.issues = issues
        super().__init__("\n".join(issue.message for issue in issues))


def number(value: object, label: str, *, allow_negative=False) -> Decimal:
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation as exc:
        raise ValueError(f"{label}: 숫자가 아닙니다 ({value})") from exc
    if not result.is_finite() or (result < 0 and not allow_negative):
        raise ValueError(f"{label}: 0 이상의 유한한 숫자가 필요합니다")
    return result


def integer(value: object, label: str) -> Decimal:
    """콤마/소수 끝전을 정리한다. float 변환 없이 Python round → int 적용."""
    result = Decimal(int(round(number(value, label, allow_negative=True))))
    if result < 0:
        raise ValueError(f"{label}: 정수화 후 0 이상의 숫자가 필요합니다")
    return result


def formula_cells(path: str | Path) -> dict[tuple[int, int], str]:
    """계산 결과가 없는 수식이 빈 값으로 사라지지 않도록 원본을 보관한다."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {(cell.row, cell.column): cell.value for row in workbook.worksheets[0].iter_rows(min_row=2)
                for cell in row if cell.data_type == "f"}
    finally:
        workbook.close()


def parse_day(value: object) -> date:
    if isinstance(value, (datetime, date)):
        day = value.date() if isinstance(value, datetime) else value
    else:
        text = str(value).strip()
        if text.replace(".0", "").isdigit() and len(text.replace(".0", "")) == 8:
            day = datetime.strptime(text.replace(".0", ""), "%Y%m%d").date()
        elif isinstance(value, (int, float, np.number)):
            day = (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(value))).date()
        else:
            day = pd.to_datetime(text, errors="raise").date()
    if not START <= day <= END:
        raise ValueError("일자는 2022년 내에 있어야 합니다")
    return day


def read_records(path: str | Path, role: str, category: str) -> list[Record]:
    frame = pd.read_excel(path, dtype=object, engine="openpyxl", keep_default_na=False)
    formulas = formula_cells(path)
    formula_rows = {row for row, _ in formulas}
    if formulas:
        # pandas가 끝부분의 미계산 수식행을 빈 행으로 잘라도 원래 행을 복원한다.
        frame = frame.reindex(range(max(len(frame), max(row for row, _ in formulas) - 1)), fill_value="")
    frame.columns = [NORMAL.get(str(c).strip(), str(c).strip()) for c in frame.columns]
    if frame.columns.duplicated().any():
        raise ValueError("같은 표준 열에 해당하는 열이 둘 이상 있습니다")
    required = {"재질", "수량", "단가", "금액"} if category == "원재료" else {"품명", "규격", "수량"}
    if category == "부재료":
        required |= {"단가", "금액"}
    if category == "제품":
        required.add("공장")
        if role != "거래":
            required.add("금액")
    if role == "거래":
        required.add("일자")
    missing = required - set(frame.columns)
    if missing:
        raise ValueError("필수 열 누락: " + ", ".join(sorted(missing)))
    records, issues = [], []
    for index, row in frame.iterrows():
        # 행 필터나 reset_index를 적용하지 않아 원본 좌측 행 번호를 보존한다.
        excel_row = int(index) + 2
        raw = {str(key): (f"{formulas[excel_row, column]} (읽은 값={value!r})"
                         if (excel_row, column) in formulas else repr(value))
               for column, (key, value) in enumerate(row.items(), 1)}
        if all(pd.isna(value) or str(value).strip() == "" for value in row) and excel_row not in formula_rows:
            continue
        source = Source(Path(path).name, file_kind(category, role), excel_row, tuple(raw.items()))
        item_parts = [str(row.get(key, "")).strip() for key in
                      (("재질",) if category == "원재료" else ("공장", "품명", "규격") if category == "제품" else ("품명", "규격"))]
        item = " / ".join(part for part in item_parts if part) or "(품목 정보 없음)"
        try:
            def cell(key: str) -> str:
                value = row.get(key)
                text = "" if pd.isna(value) else str(value).strip()
                if not text and key in frame.columns:
                    formula = formulas.get((excel_row, frame.columns.get_loc(key) + 1))
                    if formula:
                        raise ValueError(f"{key}: 수식 {formula!r}의 저장된 계산값이 없습니다. Excel에서 재계산 후 저장해 주세요")
                return text

            name = cell("재질") if category == "원재료" else cell("품명")
            spec = "" if category == "원재료" else cell("규격")
            unit = cell("단위")
            factory = cell("공장") if category == "제품" else ""
            if not name or (category == "제품" and not factory):
                raise ValueError("재질/품명 및 제품 공장은 비워둘 수 없습니다")
            qty = integer(cell("수량"), "수량")
            quantity_only = category == "제품" and role == "거래"
            price = None if quantity_only else number(cell("단가"), "단가") if cell("단가") else None
            amount = Decimal(0) if quantity_only else integer(cell("금액"), "금액") if cell("금액") else None
            if amount is None:
                if price is None:
                    raise ValueError("단가 또는 금액이 필요합니다")
                amount = Decimal(int(round(qty * price)))
            if not qty and amount:
                raise ValueError("수량이 0이면 금액도 0이어야 합니다")
            if price is None:
                price = amount / qty if qty else Decimal(0)
            day = START if role == "이월" else END if role == "재고" else parse_day(row["일자"])
            records.append(Record(name, spec, unit, day, qty, price, amount, factory, source))
        except (ValueError, TypeError, OverflowError) as exc:
            issues.append(Issue("오류", category, f"[{source.kind}] {item}",
                                f"{source.describe()} / 품목: {item} / 오류: {exc}"))
    if issues:
        raise ExcelLoadError(issues)
    return records
