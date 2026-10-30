from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.utils import get_column_letter

from .engine import verify
from .models import CATEGORIES, ZERO, Result

NUMERIC = ["입고수량", "입고단가", "입고금액", "출고수량", "출고단가", "출고금액", "재고수량", "재고단가", "재고금액"]
IDENTITY = {"원재료": ["재질"], "부재료": ["품명", "규격"], "제품": ["공장", "품명", "규격"]}
HEADERS = {c: ["일자", "구분", *IDENTITY[c], *NUMERIC] for c in CATEGORIES}
SUMMARY = ["구분", "공장", "재질 / 품명", "규격", "기초수량", "기초금액", "입고수량", "입고금액",
           "출고수량", "출고금액", "재고수량", "재고단가", "재고금액"]
THIN = Side(style="thin", color="888888")
MEDIUM = Side(style="medium", color="444444")


def prepare(ws, company: str, category: str, approval: bool, headers: list[str]):
    size = len(headers)
    last = get_column_letter(size)
    ws.sheet_view.showGridLines = False
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = "2022년도 재고수불부"
    ws["A1"].font = Font(name="맑은 고딕", size=18, bold=True, underline="single")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    for row, text in [(2, f"회사명 : {company.strip()}"), (3, "대상기간 : 2022년 01월 01일 ~ 2022년 12월 31일"),
                      (4, f"구분 : {category}")]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=size - 4)
        ws.cell(row, 1, text)
        ws.cell(row, 1).font = Font(name="맑은 고딕", size=10)
        ws.row_dimensions[row].height = 20
    if approval:
        for col, label in zip(range(size - 2, size + 1), ("담당", "검토", "승인")):
            ws.cell(2, col, label)
            ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
            for row in range(2, 5):
                ws.cell(row, col).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
                ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row, col).font = Font(name="맑은 고딕", size=10)
    for col, label in enumerate(headers, 1):
        ws.cell(6, col, label)
        width = 24 if label in ("품명", "재질", "재질 / 품명") else 18 if label == "규격" else 16 if "금액" in label else 13
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "C7"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.print_title_rows = "1:6"
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.35, bottom=0.4, header=0.15, footer=0.2)
    ws.oddFooter.center.text = "- &P / &N -"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.font = "맑은 고딕"


def finish(ws, headers):
    end, size = max(6, ws.max_row), len(headers)
    for row in ws.iter_rows(min_row=6, max_row=end, max_col=size):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10, bold=cell.row == 6)
            cell.border = Border(left=MEDIUM if cell.column == 1 else THIN,
                                 right=MEDIUM if cell.column == size else THIN,
                                 top=MEDIUM if cell.row == 6 else THIN,
                                 bottom=MEDIUM if cell.row == end else THIN)
            label = headers[cell.column - 1]
            center = cell.row == 6 or label in ("일자", "구분", "공장")
            cell.alignment = Alignment(horizontal="center" if center else "right" if cell.data_type == "n" else "left",
                                       vertical="center", wrap_text=True)
            if cell.data_type == "n":
                cell.number_format = "#,##0"
            if cell.row == 6:
                cell.fill = PatternFill("solid", fgColor="DCE6F1")
        ws.row_dimensions[row[0].row].height = 26 if row[0].row == 6 else 22
    ws.print_area = f"A1:{get_column_letter(size)}{end}"
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"


def sums(entries, kind, field):
    return sum((getattr(e.record, field) for e in entries if e.kind == kind), ZERO)


def identity(record, category):
    if category == "원재료":
        return [record.name]
    if category == "부재료":
        return [record.name, record.spec]
    return [record.factory, record.name, record.spec]


def export_workbook(result: Result, path: str | Path, company: str, approval: bool = True, progress=None):
    if not result.valid:
        raise ValueError("검증 오류를 해결한 뒤 출력해 주세요")
    wb = Workbook()
    summary = wb.active
    summary.title = "수불집계표"
    prepare(summary, company, "원재료 / 부재료 / 제품", approval, SUMMARY)
    for index, category in enumerate(CATEGORIES):
        ws = wb.create_sheet(f"{category} 수불부")
        headers = HEADERS[category]
        prepare(ws, company, category, approval, headers)
        for entries in result.ledgers.get(category, []):
            verify(entries)
            for entry in entries:
                r = entry.record
                incoming, outgoing = entry.kind == "입고", entry.kind == "출고"
                ws.append([r.day, entry.kind, *identity(r, category),
                           r.qty if incoming else None, r.price if incoming else None, r.amount if incoming else None,
                           r.qty if outgoing else None, r.price if outgoing else None, r.amount if outgoing else None,
                           entry.balance_qty, entry.balance_amount / entry.balance_qty if entry.balance_qty else ZERO,
                           entry.balance_amount])
                ws.cell(ws.max_row, 1).number_format = "yyyy-mm-dd"
            opening, closing = entries[0].record, entries[-1].record
            iq, ia = sums(entries, "입고", "qty"), sums(entries, "입고", "amount")
            oq, oa = sums(entries, "출고", "qty"), sums(entries, "출고", "amount")
            # 구분 열에는 소계 문구를 넣지 않는다.
            labels = ["[ 소계 ]", *identity(closing, category)[1:]]
            ws.append([None, None, *labels, iq, None, ia, oq, None, oa,
                       closing.qty, closing.price, closing.amount])
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="EEEEEE")
            summary.append([category, closing.factory, closing.name, closing.spec, opening.qty, opening.amount,
                            iq, ia, oq, oa, closing.qty, closing.price, closing.amount])
        finish(ws, headers)
        if progress:
            progress(90 + (index + 1) * 2, f"{category} 인쇄 시트 작성")
    if summary.max_row > 6:
        end = summary.max_row
        summary.append([None, None, "합계"] + [None] * 10)
        for col in (6, 8, 10, 13):
            summary.cell(end + 1, col, sum((summary.cell(row, col).value or ZERO for row in range(7, end + 1)), ZERO))
        for cell in summary[summary.max_row]:
            cell.fill = PatternFill("solid", fgColor="EEEEEE")
    finish(summary, SUMMARY)
    forbidden = ("역산", "추정", "자동생성", "보정")
    for ws in wb:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str) and any(word in cell.value for word in forbidden):
                    raise ValueError("회사명 또는 원본 품목 정보에 출력할 수 없는 문구가 있습니다")
    destination = Path(path).resolve()
    fd, temporary = tempfile.mkstemp(suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    try:
        wb.save(temporary)
        os.replace(temporary, destination)
        if progress:
            progress(100, "저장 완료")
    finally:
        wb.close()
        Path(temporary).unlink(missing_ok=True)
